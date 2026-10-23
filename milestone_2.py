"""
SMART EVENT PLANNING PLATFORM WITH RESOURCE COORDINATION SYSTEM - Milestone 2

Features:
- Attendee Registration & Verification
- Gmail OTP Email Verification
- Ticket Generation & Real QR Code Creation
- QR Code Email Delivery
- Live OpenCV Camera QR Attendance Scanner
- Vendor Onboarding & Event Assignment
- Consolidated M1 + M2 Reporting & Excel Sync

Configure these environment variables in the same PowerShell terminal before running:
$env:GMAIL_SENDER_EMAIL="your-sender@gmail.com"
$env:GMAIL_APP_PASSWORD="your-16-character-app-password"
"""

import os
import random
import smtplib
from email.message import EmailMessage

import cv2
import openpyxl
import qrcode

try:
    from milestone_1 import (
        events,
        venues,
        resources,
        EXCEL_FILE,
        save_m1_to_excel,
        find_event,
        find_venue,
        find_resource,
    )
except ImportError as error:
    print("Error importing milestone_1.py.")
    print("Make sure milestone_1.py and milestone_2.py are in the same folder.")
    print("Details:", error)
    raise SystemExit(1)


# Read credentials when the program starts, after PowerShell variables are set.
SENDER_EMAIL = os.getenv("GMAIL_SENDER_EMAIL", "").strip()
APP_PASSWORD = os.getenv("GMAIL_APP_PASSWORD", "").replace(" ", "")


class Attendee:
    def __init__(self, reg_id, name, email, phone, ticket_id, status="Registered"):
        self.reg_id = str(reg_id)
        self.name = str(name)
        self.email = str(email)
        self.phone = str(phone)
        self.ticket_id = str(ticket_id)
        self.status = str(status)


class Vendor:
    def __init__(self, vendor_id, name, service):
        self.id = int(vendor_id)
        self.name = str(name)
        self.service = str(service)


vendors = []
ticket_number = 1001

# Ensure M1 Event objects have the M2 collections.
for event in events:
    if not hasattr(event, "attendees"):
        event.attendees = []
    if not hasattr(event, "vendors"):
        event.vendors = []


def get_gmail_credentials():
    """Return configured sender and app password, or None if missing."""
    sender = os.getenv("GMAIL_SENDER_EMAIL", "").strip()
    password = os.getenv("GMAIL_APP_PASSWORD", "").replace(" ", "")
    if not sender or not password:
        return None
    return sender, password


def save_m2_to_excel():
    """Save M1 data plus M2 attendee, vendor, and assignment data."""
    save_m1_to_excel()

    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
    else:
        wb = openpyxl.Workbook()

    try:
        if "Attendees" in wb.sheetnames:
            ws_a = wb["Attendees"]
            if ws_a.max_row:
                ws_a.delete_rows(1, ws_a.max_row)
        else:
            ws_a = wb.create_sheet("Attendees")

        ws_a.append(
            ["Registration ID", "Name", "Email", "Phone", "Event ID", "Ticket ID", "Status"]
        )
        for event in events:
            for attendee in getattr(event, "attendees", []):
                ws_a.append([
                    attendee.reg_id,
                    attendee.name,
                    attendee.email,
                    attendee.phone,
                    event.event_id,
                    attendee.ticket_id,
                    attendee.status,
                ])

        if "Vendors" in wb.sheetnames:
            ws_v = wb["Vendors"]
            if ws_v.max_row:
                ws_v.delete_rows(1, ws_v.max_row)
        else:
            ws_v = wb.create_sheet("Vendors")

        ws_v.append(["Vendor ID", "Vendor Name", "Service"])
        for vendor in vendors:
            ws_v.append([vendor.id, vendor.name, vendor.service])

        if "Vendor Assignments" in wb.sheetnames:
            ws_va = wb["Vendor Assignments"]
            if ws_va.max_row:
                ws_va.delete_rows(1, ws_va.max_row)
        else:
            ws_va = wb.create_sheet("Vendor Assignments")

        ws_va.append(["Event ID", "Vendor ID", "Service"])
        for event in events:
            for vendor in getattr(event, "vendors", []):
                ws_va.append([event.event_id, vendor.id, vendor.service])

        wb.save(EXCEL_FILE)
    finally:
        wb.close()


def find_vendor(vendor_id):
    try:
        vendor_id = int(vendor_id)
    except (TypeError, ValueError):
        return None

    for vendor in vendors:
        if vendor.id == vendor_id:
            return vendor
    return None


def load_m2_from_excel():
    """Load attendees, vendors, and assignments from the shared M1 Excel file."""
    global ticket_number

    if not os.path.exists(EXCEL_FILE):
        return

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    try:
        # Clear M2 collections first so repeated loads don't duplicate records.
        for event in events:
            event.attendees = []
            event.vendors = []

        vendors.clear()

        if "Vendors" in wb.sheetnames:
            ws_v = wb["Vendors"]
            for row in list(ws_v.iter_rows(values_only=True))[1:]:
                if row and row[0] is not None:
                    try:
                        vendor_id = int(row[0])
                    except (TypeError, ValueError):
                        continue
                    if find_vendor(vendor_id) is None:
                        vendors.append(Vendor(vendor_id, row[1] or "", row[2] or ""))

        max_ticket_number = 1000
        if "Attendees" in wb.sheetnames:
            ws_a = wb["Attendees"]
            for row in list(ws_a.iter_rows(values_only=True))[1:]:
                if not row or row[0] is None:
                    continue
                try:
                    event_id = int(row[4])
                except (TypeError, ValueError, IndexError):
                    continue

                event = find_event(event_id)
                if event is None:
                    continue

                ticket_id = str(row[5] or "")
                attendee = Attendee(
                    row[0],
                    row[1] or "",
                    row[2] or "",
                    row[3] or "",
                    ticket_id,
                    row[6] or "Registered",
                )
                event.attendees.append(attendee)

                numeric_part = ticket_id.replace("TKT", "")
                if numeric_part.isdigit():
                    max_ticket_number = max(max_ticket_number, int(numeric_part))

        ticket_number = max_ticket_number + 1

        if "Vendor Assignments" in wb.sheetnames:
            ws_va = wb["Vendor Assignments"]
            for row in list(ws_va.iter_rows(values_only=True))[1:]:
                if not row or row[0] is None or row[1] is None:
                    continue
                try:
                    event = find_event(int(row[0]))
                    vendor = find_vendor(int(row[1]))
                except (TypeError, ValueError):
                    continue

                if event is not None and vendor is not None:
                    if vendor not in event.vendors:
                        event.vendors.append(vendor)
    finally:
        wb.close()


def send_otp(email):
    """Send OTP to the attendee's email. Never expose the OTP in the console."""
    credentials = get_gmail_credentials()
    if credentials is None:
        print("\nFailed to send OTP email.")
        print("Gmail credentials are not configured.")
        print("Set GMAIL_SENDER_EMAIL and GMAIL_APP_PASSWORD in this terminal.")
        return None

    sender_email, app_password = credentials
    otp = str(random.randint(100000, 999999))

    msg = EmailMessage()
    msg["Subject"] = "Smart Event Platform - OTP Verification"
    msg["From"] = sender_email
    msg["To"] = email
    msg.set_content(
        "Hello,\n\n"
        "Your OTP for Smart Event Platform registration is:\n\n"
        f"{otp}\n\n"
        "Please do not share this OTP with anyone.\n\n"
        "Thank you."
    )

    try:
        with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=25) as server:
            server.login(sender_email, app_password)
            server.send_message(msg)

        print(f"\nOTP sent successfully to {email}")
        return otp
    except Exception as error:
        print("\nFailed to send OTP email.")
        print("Email Error:", error)
        return None


def verify_otp(email):
    """Verify the email OTP. Stop immediately if the email cannot be sent."""
    otp = send_otp(email)
    if otp is None:
        print("Registration stopped because the OTP email could not be sent.")
        return False

    for attempt in range(3):
        entered_otp = input("Enter the OTP received in your Gmail: ").strip()
        if entered_otp == otp:
            print("Email verification successful.")
            return True
        remaining = 2 - attempt
        if remaining > 0:
            print(f"Incorrect OTP. Attempts remaining: {remaining}")

    print("Maximum OTP attempts reached. Registration cancelled.")
    return False


def generate_qr(ticket_id, reg_id, event_id):
    """Create a QR image containing ticket, registration, and event IDs."""
    qr_data = f"{ticket_id}|{reg_id}|{event_id}"
    qr_image = qrcode.make(qr_data)
    filename = f"qr_{ticket_id}.png"
    qr_image.save(filename)
    return filename


def send_ticket_email(email, name, ticket, qr_file):
    """Email the ticket and attach its QR PNG. Return True only on success."""
    credentials = get_gmail_credentials()
    if credentials is None:
        print("Ticket email was not sent: Gmail credentials are not configured.")
        return False

    sender_email, app_password = credentials
    if not os.path.isfile(qr_file):
        print(f"Ticket email was not sent: QR image not found: {qr_file}")
        return False

    msg = EmailMessage()
    msg["Subject"] = f"Your Smart Event Ticket - {ticket}"
    msg["From"] = sender_email
    msg["To"] = email
    msg.set_content(
        f"Hello {name},\n\n"
        "Your event registration is complete.\n\n"
        f"Ticket ID: {ticket}\n\n"
        "Your QR code is attached to this email. Please present it at the event entrance.\n\n"
        "Thank you."
    )

    try:
        with open(qr_file, "rb") as image_file:
            msg.add_attachment(
                image_file.read(),
                maintype="image",
                subtype="png",
                filename=os.path.basename(qr_file),
            )

        with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=25) as server:
            server.login(sender_email, app_password)
            server.send_message(msg)

        print(f"Ticket and QR code emailed successfully to {email}")
        return True
    except Exception as error:
        print("Ticket email error:", error)
        return False


def view_events():
    print("\n--- Available Events ---")
    if not events:
        print("No events available. Create an event in Milestone 1 first.")
        return

    for event in events:
        print(f"ID: {event.event_id} | Name: {event.name} | Date: {event.date}")


def register_attendee():
    global ticket_number

    if not events:
        print("No events available! Create an event in Milestone 1 first.")
        return

    view_events()
    try:
        event = find_event(int(input("Event ID: ").strip()))
    except ValueError:
        print("Invalid Event ID.")
        return

    if event is None:
        print("Event Not Found.")
        return

    reg_id = input("Registration ID: ").strip()
    name = input("Name: ").strip()
    email = input("Email: ").strip()
    phone = input("Phone: ").strip()

    if not all([reg_id, name, email, phone]):
        print("All attendee fields are required.")
        return
    if "@" not in email or "." not in email.rsplit("@", 1)[-1]:
        print("Please enter a valid email address.")
        return

    for existing in getattr(event, "attendees", []):
        if existing.email.lower() == email.lower():
            print("Already registered for this event!")
            return
        if existing.reg_id == reg_id:
            print("That Registration ID is already used for this event.")
            return

    print("\n========== EMAIL VERIFICATION ==========")
    if not verify_otp(email):
        print("Registration cancelled.")
        return

    # Generate a candidate ticket and email the QR before committing registration.
    ticket = f"TKT{ticket_number}"
    try:
        qr_file = generate_qr(ticket, reg_id, event.event_id)
    except Exception as error:
        print("Could not generate QR code:", error)
        return

    if not send_ticket_email(email, name, ticket, qr_file):
        print("Registration not completed because the ticket/QR email failed.")
        print("Fix the email error and try registering again.")
        return

    attendee = Attendee(reg_id, name, email, phone, ticket)
    if not hasattr(event, "attendees"):
        event.attendees = []
    event.attendees.append(attendee)
    ticket_number += 1

    try:
        save_m2_to_excel()
    except Exception as error:
        # Roll back in-memory registration if the Excel save fails.
        event.attendees.remove(attendee)
        ticket_number -= 1
        print("Could not save registration to Excel:", error)
        print("The ticket email may already have been sent; contact the organizer.")
        return

    print("\n========== REGISTRATION SUCCESSFUL ==========")
    print("Registration ID:", reg_id)
    print("Name           :", name)
    print("Email          :", email)
    print("Phone          :", phone)
    print("Ticket ID      :", ticket)
    print("QR Image File  :", qr_file)


def view_attendees():
    view_events()
    try:
        event = find_event(int(input("Event ID: ").strip()))
    except ValueError:
        print("Invalid Event ID.")
        return

    if event is None:
        print("Event Not Found.")
        return

    attendees_list = getattr(event, "attendees", [])
    if not attendees_list:
        print("No registrations found for this event.")
        return

    print(f"\nRegistered Participants for '{event.name}':")
    for attendee in attendees_list:
        print(
            f"Registration ID: {attendee.reg_id} | Name: {attendee.name} | "
            f"Ticket: {attendee.ticket_id} | Status: {attendee.status}"
        )


def mark_attendance():
    view_events()
    try:
        event = find_event(int(input("Event ID: ").strip()))
    except ValueError:
        print("Invalid Event ID.")
        return

    if event is None:
        print("Event Not Found.")
        return

    if not getattr(event, "attendees", []):
        print("No attendees registered for this event yet.")
        return

    print("\n====================================")
    print("  QR ATTENDANCE SCANNER (OpenCV)")
    print("  Show the participant's QR code")
    print("  Press 'q' in camera view to exit")
    print("====================================")

    camera = cv2.VideoCapture(0)
    if not camera.isOpened():
        print("Camera could not be opened.")
        camera.release()
        return

    detector = cv2.QRCodeDetector()
    last_scanned_data = None

    try:
        while True:
            ret, frame = camera.read()
            if not ret:
                print("Camera capture error.")
                break

            data, points, _ = detector.detectAndDecode(frame)
            cv2.imshow("QR Attendance Scanner", frame)

            if data and data != last_scanned_data:
                last_scanned_data = data
                print("\nScanned QR Data:", data)
                parts = data.split("|")

                if len(parts) != 3:
                    print("Invalid QR Code Format.")
                else:
                    ticket, reg_id, qr_event_id_text = parts
                    try:
                        qr_event_id = int(qr_event_id_text)
                    except ValueError:
                        print("Invalid Event ID inside QR.")
                        qr_event_id = None

                    if qr_event_id is not None and qr_event_id != event.event_id:
                        print("QR belongs to a different event.")
                    elif qr_event_id is not None:
                        matched_attendee = None
                        for attendee in getattr(event, "attendees", []):
                            if attendee.ticket_id == ticket:
                                matched_attendee = attendee
                                break

                        if matched_attendee is None:
                            print("Invalid ticket for this event.")
                        elif matched_attendee.reg_id != reg_id:
                            print("Registration verification failed.")
                        elif matched_attendee.status == "Checked In":
                            print("This attendee is already checked in.")
                        else:
                            matched_attendee.status = "Checked In"
                            try:
                                save_m2_to_excel()
                            except Exception as error:
                                matched_attendee.status = "Registered"
                                print("Could not save attendance:", error)
                            else:
                                print("\n========== ATTENDANCE SUCCESS ==========")
                                print("Name           :", matched_attendee.name)
                                print("Registration ID:", matched_attendee.reg_id)
                                print("Ticket ID      :", matched_attendee.ticket_id)
                                print("Event          :", event.name)
                                print("Status         :", matched_attendee.status)
                                break

            if cv2.waitKey(1) & 0xFF == ord("q"):
                print("Scanner closed.")
                break
    finally:
        camera.release()
        cv2.destroyAllWindows()


def add_vendor():
    try:
        vendor_id = int(input("Vendor ID (numeric): ").strip())
    except ValueError:
        print("Invalid Vendor ID.")
        return

    if find_vendor(vendor_id) is not None:
        print("Vendor ID already exists.")
        return

    name = input("Vendor Name: ").strip()
    service = input("Service Provided: ").strip()
    if not name or not service:
        print("Vendor name and service are required.")
        return

    vendors.append(Vendor(vendor_id, name, service))
    try:
        save_m2_to_excel()
    except Exception as error:
        vendors.pop()
        print("Could not save vendor:", error)
        return

    print(f"Vendor '{name}' added successfully!")


def view_vendors():
    if not vendors:
        print("No Vendors Available.")
        return

    print("\n--- Onboarded Vendors ---")
    for vendor in vendors:
        print(f"ID: {vendor.id} | Name: {vendor.name} | Service: {vendor.service}")


def assign_vendor():
    view_events()
    try:
        event = find_event(int(input("Event ID: ").strip()))
    except ValueError:
        print("Invalid Event ID.")
        return

    if event is None:
        print("Event Not Found.")
        return

    view_vendors()
    if not vendors:
        return

    try:
        vendor = find_vendor(int(input("Vendor ID to assign: ").strip()))
    except ValueError:
        print("Invalid Vendor ID.")
        return

    if vendor is None:
        print("Vendor Not Found.")
        return

    if not hasattr(event, "vendors"):
        event.vendors = []

    if vendor in event.vendors:
        print("Vendor is already assigned to this event.")
        return

    event.vendors.append(vendor)
    try:
        save_m2_to_excel()
    except Exception as error:
        event.vendors.remove(vendor)
        print("Could not save vendor assignment:", error)
        return

    print(f"Vendor '{vendor.name}' assigned to '{event.name}' successfully!")


def report():
    print("\n================ SYSTEM REPORT (M1 + M2) ================")
    if not events:
        print("No Events Found.")
        return

    for event in events:
        attendees_list = getattr(event, "attendees", [])
        vendors_list = getattr(event, "vendors", [])
        checked_in = sum(1 for attendee in attendees_list if attendee.status == "Checked In")

        print(f"\nEvent Name   : {event.name} (ID: {event.event_id})")
        print(f"Date & Time  : {event.date} | {event.start_time} - {event.end_time}")
        print(f"Registrations: {len(attendees_list)}")
        print(f"Checked In   : {checked_in}")
        print("Vendors Assigned:")

        if not vendors_list:
            print("  None")
        else:
            for vendor in vendors_list:
                print(f"  - {vendor.name} ({vendor.service})")

    print("=========================================================")


def main():
    while True:
        print("""
=============================================================
 SMART EVENT PLANNING PLATFORM - MILESTONE 2
=============================================================
1. Register Attendee
2. View Attendees
3. Mark Attendance using QR Scanner (OpenCV)
4. Add Vendor
5. View Vendors
6. Assign Vendor to Event
7. Report
8. Exit
=============================================================
""")
        choice = input("Choice (1-8): ").strip()

        if choice == "1":
            register_attendee()
        elif choice == "2":
            view_attendees()
        elif choice == "3":
            mark_attendance()
        elif choice == "4":
            add_vendor()
        elif choice == "5":
            view_vendors()
        elif choice == "6":
            assign_vendor()
        elif choice == "7":
            report()
        elif choice == "8":
            print("Exiting Milestone 2 System. Thank You!")
            break
        else:
            print("Invalid Choice! Please enter a number from 1 to 8.")


# Load saved M2 data before showing the menu.
load_m2_from_excel()


if __name__ == "__main__":
    main()
