"""
SMART EVENT PLANNING PLATFORM - MILESTONE 4
Analytics, Testing, Live Web Feedback & Finalization
Matching Slide Blueprint (Modules 01 - 07)
"""

import os
import csv
import threading
import webbrowser
from datetime import datetime
import openpyxl
from flask import Flask, render_template_string, request

EXCEL_FILE = "Smart_Event_Platform.xlsx"

# Global Datasets Across All Milestones
events = {}
expenses = {}
sponsors = {}
attendance_data = {}
resources = {}
participant_feedback = {}

# Initialize Flask App
app = Flask(__name__)

# ============================================================
# HTML TEMPLATE FOR LIVE FEEDBACK FORM
# ============================================================
HTML_FORM_TEMPLATE = """
<!DOCTYPE html>
<html>
<head>
    <title>Smart Event Platform - Live Feedback</title>
    <style>
        body { font-family: Arial, sans-serif; background-color: #f4f6f9; margin: 0; padding: 40px; }
        .container { max-width: 500px; background: white; padding: 30px; border-radius: 10px; box-shadow: 0 4px 10px rgba(0,0,0,0.1); margin: 0 auto; }
        h2 { color: #2c3e50; text-align: center; }
        .event-info { background: #eef2f7; padding: 10px 15px; border-radius: 5px; margin-bottom: 20px; font-weight: bold; }
        label { display: block; margin-top: 15px; color: #34495e; font-weight: bold; }
        select, textarea { width: 100%; padding: 10px; margin-top: 5px; border: 1px solid #ccc; border-radius: 5px; box-sizing: border-box; }
        button { width: 100%; background: #27ae60; color: white; padding: 12px; border: none; border-radius: 5px; font-size: 16px; font-weight: bold; margin-top: 20px; cursor: pointer; }
        button:hover { background: #219150; }
        .success { color: #27ae60; text-align: center; font-weight: bold; margin-top: 20px; }
    </style>
</head>
<body>
    <div class="container">
        <h2>Event Feedback Form</h2>
        {% if submitted %}
            <div class="success">
                <h3>Thank You!</h3>
                <p>Your feedback has been submitted and saved directly to the Excel database.</p>
            </div>
        {% else %}
            <div class="event-info">Event: {{ event.name }} (ID: {{ event.id }})</div>
            <form method="POST" action="/submit">
                <input type="hidden" name="event_id" value="{{ event.id }}">
                
                <label for="rating">Rating (1 to 5 Stars):</label>
                <select name="rating" id="rating" required>
                    <option value="5">⭐⭐⭐⭐⭐ (5/5 - Excellent)</option>
                    <option value="4">⭐⭐⭐⭐ (4/5 - Good)</option>
                    <option value="3">⭐⭐⭐ (3/5 - Average)</option>
                    <option value="2">⭐⭐ (2/5 - Poor)</option>
                    <option value="1">⭐ (1/5 - Very Poor)</option>
                </select>

                <label for="comment">Comments / Suggestions:</label>
                <textarea name="comment" id="comment" rows="4" placeholder="Type your experience here..." required></textarea>

                <button type="submit">Submit Feedback</button>
            </form>
        {% endif %}
    </div>
</body>
</html>
"""

# ============================================================
# EXCEL INTEGRATION & PERSISTENCE
# ============================================================

def load_all_milestone_data():
    """Reads existing records from M1, M2, M3, and M4 Excel worksheets into memory."""
    global events, expenses, sponsors, attendance_data, resources, participant_feedback

    if not os.path.exists(EXCEL_FILE):
        # Default starter datasets
        events[1] = {"id": 1, "name": "AI Workshop", "date": "2026-10-10", "venue": "Hall A", "budget": 100000.0}
        events[2] = {"id": 2, "name": "Techworld Event", "date": "2026-10-15", "venue": "Seminar Hall", "budget": 150000.0}
        attendance_data[1] = {"registered": 50, "checked_in": 42}
        attendance_data[2] = {"registered": 30, "checked_in": 18}
        resources[1] = {"event_id": 1, "name": "Projectors", "required": 5, "available": 5}
        resources[2] = {"event_id": 2, "name": "Chairs", "required": 100, "available": 70}
        return

    try:
        wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)

        # 1. Load M1 Events
        if "Events" in wb.sheetnames:
            ws = wb["Events"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None:
                    e_id = int(row[0])
                    events[e_id] = {
                        "id": e_id,
                        "name": str(row[1]) if row[1] else f"Event_{e_id}",
                        "date": str(row[2]) if row[2] else "2026-10-10",
                        "venue": str(row[5]) if len(row) > 5 and row[5] else "Not Assigned",
                        "budget": 100000.0
                    }

        # 2. Load M2 Attendance
        if "Attendees" in wb.sheetnames:
            ws = wb["Attendees"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None and len(row) > 1:
                    e_id = int(row[1]) if row[1] else 0
                    if e_id in events:
                        if e_id not in attendance_data:
                            attendance_data[e_id] = {"registered": 0, "checked_in": 0}
                        attendance_data[e_id]["registered"] += 1
                        if len(row) > 6 and str(row[6]).strip().lower() in ["checked in", "present"]:
                            attendance_data[e_id]["checked_in"] += 1

        # 3. Load M3 Expenses
        if "Expenses" in wb.sheetnames:
            ws = wb["Expenses"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None:
                    exp_id = int(row[0])
                    e_id = int(row[1]) if row[1] else 0
                    expenses[exp_id] = {
                        "id": exp_id,
                        "event_id": e_id,
                        "description": str(row[3]) if len(row) > 3 and row[3] else "Expense",
                        "amount": float(row[4]) if len(row) > 4 and row[4] else 0.0,
                        "status": "Approved"
                    }

        # 4. Load M3 Sponsors
        if "Sponsors" in wb.sheetnames:
            ws = wb["Sponsors"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None:
                    sp_id = int(row[0])
                    e_id = int(row[1]) if row[1] else 0
                    sponsors[sp_id] = {
                        "id": sp_id,
                        "event_id": e_id,
                        "amount": float(row[3]) if len(row) > 3 and row[3] else 0.0,
                        "status": str(row[4]) if len(row) > 4 and row[4] else "Confirmed"
                    }

        # 5. Load M4 Feedback Records
        if "Feedback" in wb.sheetnames:
            ws = wb["Feedback"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None and len(row) > 1:
                    e_id = int(row[1]) if row[1] else 0
                    if e_id in events:
                        rating = float(row[2]) if row[2] is not None else 5.0
                        comment = str(row[3]) if len(row) > 3 and row[3] else ""
                        participant_feedback.setdefault(e_id, []).append({
                            "rating": rating,
                            "comment": comment
                        })

        wb.close()
    except Exception as err:
        print(f"⚠️ [Sync Warning] Excel read error: {err}")


def save_feedback_to_excel(event_id, rating, comment):
    """Saves live feedback directly into Smart_Event_Platform.xlsx."""
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        wb.save(EXCEL_FILE)

    wb = openpyxl.load_workbook(EXCEL_FILE)
    if "Feedback" not in wb.sheetnames:
        ws = wb.create_sheet(title="Feedback")
        ws.append(["Feedback ID", "Event ID", "Rating (1-5)", "Comments", "Submission Timestamp"])
    else:
        ws = wb["Feedback"]

    feedback_id = ws.max_row
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws.append([feedback_id, event_id, rating, comment, timestamp])
    wb.save(EXCEL_FILE)
    wb.close()
    print(f"\n[LIVE SYNC] Saved Feedback -> Event ID: {event_id} | Rating: {rating} | Comment: {comment}")

# ============================================================
# FLASK LIVE WEB ROUTES
# ============================================================

@app.route('/feedback/<int:event_id>')
def feedback_page(event_id):
    if event_id not in events:
        return "Event Not Found", 404
    return render_template_string(HTML_FORM_TEMPLATE, event=events[event_id], submitted=False)

@app.route('/submit', methods=['POST'])
def submit_feedback():
    event_id = int(request.form['event_id'])
    rating = float(request.form['rating'])
    comment = request.form['comment'].strip()

    participant_feedback.setdefault(event_id, []).append({"rating": rating, "comment": comment})
    save_feedback_to_excel(event_id, rating, comment)

    event = events.get(event_id, {"name": "Event", "id": event_id})
    return render_template_string(HTML_FORM_TEMPLATE, event=event, submitted=True)

def run_server():
    app.run(port=5000, debug=False, use_reloader=False)

# ============================================================
# MILESTONE 4 BLUEPRINT MODULES (01 - 07)
# ============================================================

def money(amount):
    return f"Rs. {amount:,.2f}"

def total_expenses(event_id):
    return sum(exp["amount"] for exp in expenses.values() if exp["event_id"] == event_id and exp["status"] != "Rejected")

def total_sponsorship(event_id):
    return sum(sp["amount"] for sp in sponsors.values() if sp["event_id"] == event_id and sp["status"] == "Confirmed")


# MODULE 01: ATTENDANCE ANALYSIS
def show_attendance_analysis():
    print("\n" + "=" * 65)
    print(" 01. ATTENDANCE ANALYSIS")
    print("=" * 65)
    if not events:
        print("No event records found.")
        return

    for e_id, e in events.items():
        att = attendance_data.get(e_id, {"registered": 0, "checked_in": 0})
        reg = att["registered"]
        chk = att["checked_in"]
        rate = (chk / reg * 100) if reg > 0 else 0.0
        print(f"• Event: {e['name']} (ID: {e_id})")
        print(f"  - Registered : {reg}")
        print(f"  - Checked In : {chk}")
        print(f"  - Turnout    : {rate:.2f}%\n")


# MODULE 02: BUDGET OVERVIEW
def show_budget_overview():
    print("\n" + "=" * 65)
    print(" 02. BUDGET OVERVIEW")
    print("=" * 65)
    if not events:
        print("No event records found.")
        return

    for e_id, e in events.items():
        budget = e["budget"]
        spent = total_expenses(e_id)
        spon = total_sponsorship(e_id)
        balance = (budget + spon) - spent

        print(f"• Event: {e['name']} (ID: {e_id})")
        print(f"  - Allocated Budget   : {money(budget)}")
        print(f"  - Sponsorship Raised : {money(spon)}")
        print(f"  - Actual Expenses    : {money(spent)}")
        print(f"  - Net Balance        : {money(balance)}")
        print(f"  - Status             : {'UNDER BUDGET' if balance >= 0 else 'OVER BUDGET'}\n")


# MODULE 03: EVENT ANALYTICS (KPIs & SCORING)
def calculate_performance_score(e_id):
    att = attendance_data.get(e_id, {"registered": 0, "checked_in": 0})
    reg, chk = att["registered"], att["checked_in"]
    att_score = (chk / reg * 40) if reg > 0 else 30.0

    budget = events[e_id]["budget"]
    spent = total_expenses(e_id)
    budget_score = 20.0 if spent <= budget else 10.0

    res_score = 20.0

    fb_list = participant_feedback.get(e_id, [])
    if fb_list:
        avg_fb = sum(f["rating"] for f in fb_list) / len(fb_list)
        fb_score = (avg_fb / 5) * 20
    else:
        fb_score = 15.0

    score = round(att_score + budget_score + res_score + fb_score, 2)
    grade = "EXCELLENT" if score >= 80 else "GOOD" if score >= 60 else "AVERAGE"
    return score, grade

def show_event_analytics():
    print("\n" + "=" * 65)
    print(" 03. EVENT ANALYTICS & PERFORMANCE KPIS")
    print("=" * 65)
    for e_id, e in events.items():
        score, grade = calculate_performance_score(e_id)
        fb_list = participant_feedback.get(e_id, [])
        avg_fb = f"{sum(f['rating'] for f in fb_list)/len(fb_list):.1f}/5" if fb_list else "N/A"
        print(f"• Event ID {e_id} ({e['name']}): Score = {score}/100 | Grade = {grade} | Avg Feedback = {avg_fb}")


# MODULE 04: RESOURCE REQUIREMENTS & SHORTAGES
def show_resource_requirements():
    print("\n" + "=" * 65)
    print(" 04. RESOURCE REQUIREMENTS & SHORTAGE ANALYSIS")
    print("=" * 65)
    if not resources:
        print("No resource allocations recorded.")
        return

    for r_id, r in resources.items():
        e_id = r["event_id"]
        req, avail = r["required"], r["available"]
        shortage = req - avail if req > avail else 0

        print(f"• Event ID {e_id} | Resource: {r['name']}")
        print(f"  - Required: {req} | Available: {avail}")
        if shortage > 0:
            print(f"  - ❌ ALERT: Shortage of {shortage} units detected!")
        else:
            print(f"  - ✅ Resource levels optimal.")


# MODULE 05 & 06: REPORTS & CSV EXPORT
def export_csv_report():
    print("\n" + "=" * 65)
    print(" 05 & 06. GENERATING CONSOLIDATED CSV REPORT")
    print("=" * 65)
    filename = "final_event_report.csv"
    try:
        with open(filename, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Event ID", "Event Name", "Date", "Venue", "Registered", "Checked In", "Budget", "Expenses", "Sponsorship", "Performance Score", "Grade"])

            for e_id, e in events.items():
                att = attendance_data.get(e_id, {"registered": 0, "checked_in": 0})
                spent = total_expenses(e_id)
                spon = total_sponsorship(e_id)
                score, grade = calculate_performance_score(e_id)

                writer.writerow([e_id, e["name"], e["date"], e["venue"], att["registered"], att["checked_in"], e["budget"], spent, spon, score, grade])

        print(f"✅ Success: Final consolidated report exported to '{filename}'!")
    except Exception as err:
        print(f"❌ Failed to export CSV report: {err}")


# MODULE 07: FINAL SYSTEM DASHBOARD
def show_final_dashboard():
    print("\n" + "=" * 70)
    print(" 07. FINAL EXECUTIVE SYSTEM DASHBOARD")
    print("=" * 70)
    load_all_milestone_data()

    show_attendance_analysis()
    show_budget_overview()
    show_event_analytics()
    show_resource_requirements()


# ============================================================
# CLI MAIN DASHBOARD
# ============================================================

def milestone4_menu():
    load_all_milestone_data()

    # Start live feedback server in background thread
    server_thread = threading.Thread(target=run_server, daemon=True)
    server_thread.start()

    while True:
        print("\n" + "=" * 65)
        print("    MILESTONE 4 - ANALYTICS, TESTING & FINALIZATION")
        print("=" * 65)
        print("1. Attendance Analysis (01)")
        print("2. Budget Overview (02)")
        print("3. Event Analytics & Performance KPIs (03)")
        print("4. Resource Requirements & Shortages (04)")
        print("5. Launch Live Web Feedback Server")
        print("6. Export Final Consolidated CSV Report (05 & 06)")
        print("7. View Final Executive System Dashboard (07)")
        print("8. Exit Platform")

        choice = input("\nEnter choice (1-8): ").strip()

        if choice == "1":
            show_attendance_analysis()
        elif choice == "2":
            show_budget_overview()
        elif choice == "3":
            show_event_analytics()
        elif choice == "4":
            show_resource_requirements()
        elif choice == "5":
            for e_id, e in events.items():
                print(f"ID: {e_id} | Name: {e['name']}")
            try:
                e_id = int(input("\nEnter Event ID: ").strip())
                if e_id in events:
                    url = f"http://127.0.0.1:5000/feedback/{e_id}"
                    print(f"\n🔗 Opening Live Feedback Server: {url}")
                    webbrowser.open(url)
                else:
                    print("Event not found.")
            except ValueError:
                print("Invalid ID.")
        elif choice == "6":
            export_csv_report()
        elif choice == "7":
            show_final_dashboard()
        elif choice == "8":
            print("Exiting Milestone 4. Goodbye!")
            break

if __name__ == "__main__":
    milestone4_menu()